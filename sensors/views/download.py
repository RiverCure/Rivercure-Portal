from django.shortcuts import get_object_or_404
from sensors.models import Sensor, SensorClass, SensorClassProperty, SensorObservation
from organization.models import Membership
from openpyxl import Workbook
from django.conf import settings
import os
from datetime import datetime as d
from django.http import Http404, HttpResponse
import string
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

LAST_INDEX = 1048576

def readme(ws, date_str, user):
    font = Font(size=20)
    ws.title = 'README'
    ws['A1'] = f'This is a auto-generated Excel file produced by the Rivercure Portal system at {date_str} for user {user}'
    ws['A1'].font = Font(size=18)
    ws['A2'] = 'This file is tailored to the user that generated it, so its submission might not work in two different user accounts'
    ws['A2'].font = Font(size=18)
    ws['A3'] = 'All rights reserved to Rivercure Project'
    ws['A3'].font = Font(size=10)
    ws.protection.sheet = True

def increase_column_width(ws):
    # Increase column width
        dims = {}
        for row in ws.rows:
            for cell in row:
                if cell.value:
                    dims[cell.column_letter] = max((dims.get(cell.column_letter, 0), len(str(cell.value))))
        for col, value in dims.items():
            ws.column_dimensions[col].width = str(int(value) + 7)

# Download file
def download_file(wb, file_name):
    relative_path = os.path.join(settings.MEDIA_ROOT, 'excels')
    if not os.path.exists(relative_path):
        os.makedirs(relative_path)
    full_path = os.path.join(relative_path, file_name)
    print(full_path)
    print('Hey1')
    try:
        wb.save(full_path)
    except Exception as e:
        print(e)
    print('Hey2')
    if os.path.exists(full_path):
        with open(full_path, 'rb') as fh:
            response = HttpResponse(fh.read(), content_type="application/vnd.ms-excel")
            response['Content-Disposition'] = 'inline; filename=' + os.path.basename(full_path)
            return response
    raise Http404

def sensor_excel_download(request):
    date = d.now()
    date_str = date.strftime('%Y-%m-%d_%H:%M:%S')

    wb = Workbook()
    
    def sensors():
        # Same query as in sensor-list
        ws = wb.create_sheet('Sensors')
        sensorClasses = SensorClass.objects.filter(state='active', organization__membership__in=Membership.objects.filter(user=request.user, access_granted=True))
        sensorClassCodes = list()
        for sensorClass in sensorClasses:
            sensorClassCodes.append(sensorClass.code)

        # Colors
        redFill = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid') # for mandatory fields
        yellowFill = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid') # for optional fields
        
        # Validators
        validator_title = 'Invalid Entry'
        yes_no_v = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        yes_no_v.error ='Please select one of the options (Yes or No) or leave it blank for default value'
        yes_no_v.errorTitle = validator_title
        ws.add_data_validation(yes_no_v)

        sensor_class_v = DataValidation(type="list", formula1=f'"{",".join(sensorClassCodes)}"', allow_blank=False)
        sensor_class_v.error =f'Please select one of the options ({", ".join(sensorClassCodes)})'
        sensor_class_v.errorTitle = validator_title
        ws.add_data_validation(sensor_class_v)

        states = Sensor._meta.get_field('state').choices
        states_full_name = [ state[1] for state in states ]
        state_v = DataValidation(type="list", formula1=f'"{",".join(states_full_name)}"', allow_blank=False)
        state_v.error =f'Please select one of the options ({", ".join(states_full_name)})'
        state_v.errorTitle = validator_title
        ws.add_data_validation(state_v)

        ws['A1'] = 'Sensor class*'
        ws['A1'].fill = redFill
        sensor_class_v.add(f'A2:A{LAST_INDEX}')
        ws['B1'] = 'Code*'
        ws['B1'].fill = redFill
        ws['C1'] = 'Name*'
        ws['C1'].fill = redFill
        ws['D1'] = 'State*'
        ws['D1'].fill = redFill
        state_v.add(f'D2:D{LAST_INDEX}')
        ws['E1'] = 'Description'
        ws['E1'].fill = yellowFill
        ws['F1'] = 'Is public'
        ws['F1'].fill = yellowFill
        yes_no_v.add(f'F2:F{LAST_INDEX}')
        ws['G1'] = 'Local'
        ws['G1'].fill = redFill
        
        # Optional/mandatory indicators
        ws['H1'].fill = redFill
        ws['I1'] = 'Required field'
        ws['K1'].fill = yellowFill
        ws['L1'] = 'Optional field'

        increase_column_width(ws)
    
    readme(wb.active, date_str, request.user)
    sensors()
    return download_file(wb, f'RCP_sensors_create_{request.user}_{date_str}.xlsx')

def sensor_observations_excel_download(request, sensorId):
    sensor = get_object_or_404(Sensor, pk=sensorId)
    sensorClass = sensor.sensorClass
    properties = SensorClassProperty.objects.filter(sensorClass=sensorClass).exclude(derivedBy__isnull=False).order_by('name')

    date = d.now()
    date_str = date.strftime('%Y-%m-%d_%H:%M:%S')

    wb = Workbook()

    def prop_comment(prop):
        if prop.type == 'number':
            comment_str = 'This property is of type number, '
            comment_str += '' if prop.thresholdLowerNoncritical is None else f'has threshold lower non-critical of {prop.thresholdLowerNoncritical}, '
            comment_str += '' if prop.thresholdUpperNoncritical is None else f'threshold upper non-critical of {prop.thresholdUpperNoncritical}'
            comment_str += '' if prop.thresholdLowerCritical is None else f'threshold lower critical of {prop.thresholdLowerCritical}'
            comment_str += '' if prop.thresholdUpperCritical is None else f'threshold upper critical of {prop.thresholdUpperCritical}'
            comment_str += 'and the unit is '
            comment_str += 'not specified' if prop.unit is None else prop.unit
            return Comment(comment_str, 'Rivercure Portal')
        else:
            comment_str = f'This property is of type {prop.type}'
            return Comment(comment_str, 'Rivercure Portal')

    def observations():
        ws = wb.create_sheet('Observations')

        # Colors
        redFill = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid') # for mandatory fields
        yellowFill = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid') # for optional fields

        # Validators
        validator_title = 'Invalid Entry'
        choices = SensorObservation._meta.get_field('severity').choices
        choices_full_name = [ choice[1] for choice in choices ]
        severity_v = DataValidation(type='list', formula1=f'"{",".join(choices_full_name)}"', allow_blank=True)
        severity_v.error =f'Please select one of the options ({", ".join(choices_full_name)})'
        severity_v.errorTitle = validator_title
        ws.add_data_validation(severity_v)

        date_v = DataValidation(type='date')
        date_v.error = 'Please introuce a valid date'
        date_v.errorTitle = validator_title
        ws.add_data_validation(date_v)

        time_v = DataValidation(type='time')
        time_v.error = 'Please introuce a valid time'
        time_v.errorTitle = validator_title
        ws.add_data_validation(time_v)

        ws['A1'] = 'Date'
        ws['A1'].fill = redFill
        date_v.add(f'A2:A{LAST_INDEX}')
        ws['B1'] = 'Time'
        ws['B1'].fill = redFill
        time_v.add(f'B2:B{LAST_INDEX}')
        ws['C1'] = 'Severity'
        ws['C1'].fill = redFill
        severity_v.add(f'C2:C{LAST_INDEX}')

        # Uppercase list of letters of the alphabet, starting in D
        alphabet = list(string.ascii_uppercase)[3:]

        last_idx = 0
        for idx, prop in enumerate(properties):
            cell = f'{alphabet[idx]}1'
            ws[cell] = f'{prop.name} value'
            ws[cell].fill = yellowFill if prop.isOptional else redFill
            ws[cell].comment = prop_comment(prop)
            last_idx = idx
        
        # Optional/mandatory indicators
        offset = 3
        ws[f'{alphabet[last_idx + offset]}1'].fill = redFill
        ws[f'{alphabet[last_idx + offset + 1]}1'] = 'Required field'
        ws[f'{alphabet[last_idx + offset + 2]}1'].fill = yellowFill
        ws[f'{alphabet[last_idx + offset + 2 + 1]}1'] = 'Optional field'

        increase_column_width(ws)

    readme(wb.active, date_str, request.user)
    observations()
    return download_file(wb, f'RCP_observations_{sensor.name}_create_{request.user}_{date_str}.xlsx')